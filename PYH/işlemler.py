# -*- coding: utf-8 -*-
"""
Created on Sun Nov  6 12:35:36 2022

@author: Ruzgar
"""
import random
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

Rapor = Workbook()
ws = Rapor.active
ws.title = "İşlemler"

#Data
Başlıklar = ["Ad-Soyad", "İşlem Türü", "Şehir"]
Kişiler = [
    ["Fofofo","","İstanbul"],
    ["Lololo","","İstanbul"],
    ["Rororo","","İstanbul"],
    ["Kokoko","","İstanbul"],
    ["Şoşoşo","","İstanbul"],
    ["Vovovo","","İstanbul"]
    ]

#İşlem başarısı belirleme 
def İşlem(Data):
    Data = Kişiler
    A = "Başarılı"
    B = "Başarısız"
    for k in Data:
       k[1] = random.choice([A,B])
       
İşlem(Kişiler)

#Dosyaya yazma
ws.append(Başlıklar)
for row in Kişiler:
    ws.append(row)

#Tablo oluşturma
reff="C"+str(len(Kişiler)+1)
tab = Table(displayName="BaşarıOranıTablosu", ref="A1:"+reff)
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)



#İstatistik
İşlemSayısı = ws.max_row - 1
Data = Kişiler
BaşarıSayısı=0
BşarısızSayısı=0
for k in Data:
    if (k[1]=="Başarılı"):
        BaşarıSayısı+=1
    else:
        BşarısızSayısı+=1

BaşarıYüzdesi = BaşarıSayısı/İşlemSayısı*100

#İstatistik Tablosu 
Başlık = ["İşlem Sayısı", "Başarılı İşlem", "Başarısız İşlem", "Başarı Yüzdesi"]
Değerler = [İşlemSayısı,BaşarıSayısı,BşarısızSayısı,BaşarıYüzdesi]

for i in range(0,4):
    print(Başlık[i] + " : " + str(Değerler[i]))


ws.append([])
ws.append(Başlık)
ws.append(Değerler)


reff_s="A"+str(len(Kişiler)+3)
reff_f="D"+str(len(Kişiler)+4)
itab = Table(displayName="İstatistik", ref=reff_s+":"+reff_f)
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
itab.tableStyleInfo = style    
ws.add_table(itab)

Rapor.save("Rapor.xlsx")
